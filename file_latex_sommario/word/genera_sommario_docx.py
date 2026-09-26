"""Genera il sommario in Word con la stessa impaginazione di main.tex.

Impaginazione ripresa da file_latex_sommario/main.tex:
  - A4, margini 3 cm sopra/sotto/destra, 3,5 cm a sinistra;
  - Times New Roman 12 pt, interlinea 1,15, nessun rientro, 3 pt fra i paragrafi;
  - intestazione con logo UCBM a sinistra e riga in corsivo a destra, filetto sotto;
  - numero di pagina centrato a piè di pagina;
  - blocco del titolo: titolo in grassetto centrato, tabella con laureando, relatore,
    correlatore e sessione, filetto;
  - sezioni numerate in grassetto 14 pt, sottosezioni in grassetto 12 pt;
  - bibliografia su due colonne in corpo 6,4 pt.

Il testo si scrive in CONTENUTO con un piccolo markup:
  [i]...[/i] corsivo (variabili matematiche), [sup]...[/sup] apice,
  [sub]...[/sub] pedice.

Uso:
    python genera_sommario_docx.py
"""
import difflib
import os
import re
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

QUI = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(QUI, '..', 'figure', 'ucbm-logo.png')
USCITA = os.path.join(QUI, 'Sommario_Tesi_Dragotta.docx')
# Copia di lavoro condivisa su OneDrive: Claudio la tiene aperta e vede le modifiche.
# La copia nel repository resta come storico; e' anche il riferimento per capire se la
# copia OneDrive e' stata modificata a mano dall'ultima generazione.
ONEDRIVE = r'C:\Users\ludov\Documents\OneDrive\Documenti\Sommario_Tesi_Dragotta.docx'

FONT = 'Times New Roman'
TITOLO = ('Shor’s Algorithm under Noise: Machine-Learning Ablation, '
          'Quantum Error Correction, and Approximate QFT')
INTESTAZIONE = ('Sommario del Lavoro di Tesi, Laurea Magistrale in Ingegneria dei '
                'Sistemi Intelligenti, Università Campus Bio-Medico di Roma')
BLOCCO = [('Laureando', 'Claudio Dragotta', 'Relatore', 'Prof. Paolo Soda'),
          ('Correlatore', 'Ing. Floriano Caprio', 'Sessione di Laurea',
           'Ottobre — A.A. 2025/2026')]

# ---------------------------------------------------------------------------
# Contenuto: (livello, titolo, [paragrafi]). Livello 1 = sezione, 2 = sottosezione.
# Le sezioni ancora da trascrivere hanno l'elenco dei paragrafi vuoto.
# ---------------------------------------------------------------------------
CONTENUTO = [
    (1, 'Introduzione e obiettivi', [
        'L’algoritmo di Shor consente di fattorizzare numeri interi in tempo polinomiale '
        'rispetto alla lunghezza binaria dell’input, con implicazioni per la sicurezza dei '
        'sistemi crittografici fondati sulla difficoltà della fattorizzazione [1–3]. '
        'La realizzazione di questo vantaggio teorico richiede tuttavia un controllo degli '
        'errori compatibile con la profondità dei circuiti. Errori nelle operazioni, '
        'decoerenza e lettura imperfetta possono alterare la distribuzione delle misure e '
        'compromettere l’informazione necessaria alla ricerca dell’ordine '
        'moltiplicativo [4].',
        'La tesi sviluppa uno studio simulativo volto a stabilire quali interventi migliorino '
        'effettivamente l’affidabilità della computazione e rispetto a quali riferimenti. '
        'Il punto di partenza è la misura della sensibilità di Shor agli errori di '
        'porta; ne seguono tre strategie: la selezione degli esiti mediante post-processing, '
        'con uno studio di ablazione del contributo del machine learning; la correzione '
        'quantistica degli errori, con surface code e codici qLDPC; la riduzione del numero di '
        'porte mediante trasformata di Fourier quantistica approssimata, o AQFT.',
        'Gli ambiti di intervento — selezione degli esiti, protezione dell’informazione '
        'e riduzione del costo circuitale — sono collegati da un criterio comune: confrontare '
        'ogni soluzione con un riferimento esplicito, isolare il contributo dei singoli '
        'componenti e misurare il beneficio mediante indicatori coerenti con il compito. Le '
        'campagne sperimentali sono state condotte separatamente, difatti i metodi di correzione '
        'degli errori non sono stati applicati direttamente all’intero circuito di Shor in '
        'un’unica simulazione.',
    ]),
    (1, 'Materiali e metodi', []),
    (2, 'Algoritmo di Shor, validazione e definizione del successo', [
        'Scelta una base [i]a[/i] coprima con il numero [i]N[/i] da fattorizzare, la successione '
        '[i]a[/i][sup][i]x[/i][/sup] mod [i]N[/i] è periodica; il periodo, detto ordine '
        '[i]r[/i], è il minimo intero positivo per cui [i]a[/i][sup][i]r[/i][/sup] ≡ 1 '
        '(mod [i]N[/i]). La parte quantistica lo stima con stima di fase e QFT inversa; la parte '
        'classica ricava i fattori per approssimazione razionale e massimo comune divisore, e li '
        'verifica.',
        'L’istanza principale è [i]N[/i] = 15, [i]a[/i] = 7, con otto qubit di controllo '
        'e quattro di lavoro; l’ordine è [i]r[/i] = 4 e gli esiti ideali del registro di '
        'controllo sono 0, 64, 128 e 192. Ogni esecuzione con misura, detta [i]shot[/i], '
        'restituisce un solo esito; molti shot formano l’istogramma su cui opera il '
        'post-processing. Il successo è il recupero di fattori non banali, non la '
        'ricostruzione esatta dell’ordine. Vi conducono tre dei quattro picchi ideali, con '
        'probabilità complessiva del 75%. La verifica classica è però permissiva e '
        'accetta 63 dei 256 valori, tanto che esiti del tutto uniformi darebbero comunque un '
        'successo di 63/256 ≃ 24,61%. Questo riferimento dipende dall’istanza e dal '
        'post-processing e non misura la fedeltà dello stato quantistico. Mostra inoltre che '
        'fattorizzazioni corrette possono essere ottenute anche quando la distribuzione degli '
        'esiti di misura del registro di controllo ha perso gran parte della struttura ideale '
        'dei picchi [5].',
        'La correttezza dei circuiti è verificata indipendentemente dal successo della '
        'fattorizzazione, controllando operatori modulari, ripristino dei qubit ausiliari e '
        'accordo con le distribuzioni teoriche. Oltre a [i]N[/i] = 15, si considerano le istanze '
        '[i]N[/i] = 21, [i]a[/i] = 2 e [i]N[/i] = 35, [i]a[/i] = 6, che, validate idealmente, '
        'usano l’aritmetica reversibile di Beauregard [6]; solo [i]N[/i] = 21 entra nelle '
        'prove rumorose esplorative sull’AQFT, perché per [i]N[/i] = 35 le risorse '
        'hardware disponibili non consentivano la simulazione di un circuito di Shor adeguato. '
        'Si impiegano Qiskit/Aer, scikit-learn, Stim, PyMatching e ldpc, registrando versioni, '
        'parametri, semi casuali e identificativi dei circuiti.',
    ]),
    (2, 'Sensibilità di Shor agli errori di porta', [
        'Il punto di partenza dello studio è stato misurare quanto l’algoritmo di '
        'Shor sia robusto agli errori di porta. Nel circuito dell’istanza principale, '
        'compilato nelle porte native RZ, SX, X e CX (profondità 412, con 224 porte CX e 70 '
        'SX), si introducono errori di Pauli indipendenti dopo ogni porta, escluse le rotazioni '
        'RZ, che sono virtuali. L’intensità del rumore è fissata dalla '
        'probabilità [i]p[/i][sub]g[/sub] che dopo una porta si verifichi un errore di Pauli '
        'diverso dall’identità, fatta variare su 13 valori fra 0 e 0,5.',
        'Il modello di rumore è fenomenologico e uniforme e non riproduce la calibrazione di '
        'un dispositivo reale. Per questo [i]p[/i][sub]g[/sub] va inteso come probabilità '
        'di errore per singola porta: non è direttamente confrontabile con la '
        'probabilità di fallimento logico [i]p[/i][sub]L[/sub] dei codici di correzione, '
        'perché l’algoritmo di Shor non viene compilato in forma fault-tolerant.',
    ]),
    (2, 'Mitigazione degli errori mediante post-processing e machine learning', [
        'La prima strategia agisce a valle della misura: senza modificare il circuito, cerca di '
        'ricavare i fattori dagli esiti rumorosi scegliendo meglio quali verificare, e valuta '
        'se un modello di machine learning renda questa scelta più efficace. Tre criteri '
        'di selezione, con estrattore dei fattori e regola dei pareggi comuni, operano sugli '
        'stessi istogrammi da 1.024 shot; ogni istogramma corrisponde a un’iterazione e, '
        'se non si ottengono i fattori, se ne esegue una nuova. TOP-1 '
        '(Metodo 1) verifica il solo esito più frequente; TOP-4 fino a quattro candidati '
        'ordinati per frequenza; il Metodo 2 applica TOP-4 solo se un classificatore predice il '
        'successo del candidato dominante. Il confronto fra Metodo 2 e TOP-4 senza filtro è '
        'lo studio di ablazione, che stabilisce se il vantaggio venga dal modello appreso o dal '
        'semplice provare più candidati.',
        'Due scenari illustrativi, UC1 di riferimento e UC2 di stress, includono '
        'depolarizzazione, rilassamento termico ed errori di lettura; per ciascuno scenario, '
        '2.000 istogrammi con parametri variati entro ±50% del nominale sono divisi '
        '60%/20%/20% fra addestramento, selezione e verifica. Random forest, support vector '
        'machine (SVM) e perceptron multistrato sono confrontati su F1; il modello scelto è '
        'valutato su 400 istogrammi indipendenti per scenario. La metrica primaria è il '
        'numero di iterazioni per ottenere i fattori; l’accuratezza è stata considerata '
        'un elemento secondario nello studio del classificatore, mentre il tempo di calcolo '
        'è stato rapportato alle risorse hardware disponibili. Si usano 30 repliche appaiate '
        'per scenario e al massimo 50 iterazioni.',
    ]),
    # Numerazione dei riferimenti in ordine di prima citazione: con questo ordine dei
    # paragrafi AlphaQubit diventa [10], Bravyi et al. [11], Roffe et al. [12].
    (2, 'Correzione quantistica degli errori: surface code e codici qLDPC', [
        'La seconda strategia protegge l’informazione prima che il rumore la distrugga. '
        'Nella correzione quantistica degli errori, l’informazione di un qubit logico viene '
        'distribuita su più qubit fisici; le sindromi permettono di rilevare gli errori '
        'senza misurare direttamente lo stato logico e guidano il decoder nella scelta della '
        'correzione. Le prestazioni sono misurate attraverso la probabilità di fallimento '
        'logico [i]p[/i][sub]L[/sub].',
        'Codice a ripetizione e codice di Steane sono utilizzati come verifiche preliminari, '
        'mentre l’analisi principale riguarda circuiti di memoria basati su surface code '
        '[7–9], di distanza [i]d[/i] crescente, che fissa quanti errori il codice può '
        'correggere, con errori nelle operazioni, nella preparazione e nella misura e, in alcune '
        'configurazioni, crosstalk tra qubit adiacenti. Il decoder di riferimento è il '
        '[i]minimum-weight perfect matching[/i] (MWPM), che associa a coppie le sindromi '
        'attivate scegliendo la configurazione di errori più probabile. MWPM è '
        'confrontato con una rete neurale autonoma e con una strategia ibrida, in cui una rete '
        'neurale decide quando correggere la scelta di MWPM; la soglia d’intervento è '
        'selezionata sul validation set e mantenuta fissa sul test set [10].',
        'Uno studio esplorativo confronta inoltre il surface code con il Gross code di IBM, un '
        'codice quantum Low-Density Parity-Check (qLDPC) che codifica 12 qubit logici in 144 '
        'qubit di dato [11], nel modello code-capacity, con errori sui soli qubit di dato e '
        'sindromi perfette. BP+OSD ([i]belief propagation with ordered statistics '
        'decoding[/i]) è usato come ulteriore decoder di confronto sul surface code e come '
        'decoder del Gross code [12]; per quest’ultimo se ne confrontano più '
        'configurazioni, fra cui l’ordine di aggiornamento dei messaggi, parallelo o '
        'serial.',
        'Il confronto mostra che i due approcci non sono equivalenti: a parità di qubit '
        'logici il Gross code richiede molti meno qubit di dato, con un vantaggio che dipende '
        'in modo decisivo dal decoder (Sezione 3.3).',
    ]),
    (2, 'Riduzione del rumore con la AQFT', [
        'La terza strategia riduce l’effetto del rumore alla fonte, diminuendo il numero di '
        'porte su cui può agire. La QFT approssimata (AQFT) elimina le rotazioni controllate '
        'di piccolo angolo della QFT, che contribuiscono poco all’informazione di fase: il '
        'circuito perde un po’ di precisione, ma ha meno porte e minore profondità, e '
        'quindi meno occasioni di errore [13]. La domanda è se il rumore evitato compensi '
        'la precisione persa.',
        'L’AQFT si applica separatamente alla QFT inversa finale di Shor per [i]N[/i] = 15 '
        'e alle QFT interne all’aritmetica modulare per [i]N[/i] = 21; benchmark distinti '
        'di stima di fase, con registri da sei a dieci qubit, isolano l’effetto '
        'dell’AQFT. Il modello di rumore è costruito a partire dalle infedeltà '
        'delle porte e dagli errori di lettura della calibrazione offline FakeSherbrooke, senza '
        'aggiungere contributi [i]T[/i][sub]1[/sub]/[i]T[/i][sub]2[/sub], mantenendo invariati '
        'layout e semi casuali. Il grado di approssimazione dell’AQFT, cioè quante '
        'rotazioni si eliminano, viene selezionato su un insieme dedicato e valutato su dati di '
        'verifica indipendenti; gli shot per configurazione, 8.192 nel pilota per [i]N[/i] = 15 '
        'e 256 per [i]N[/i] = 21, sono suddivisi equamente tra selezione e verifica.',
    ]),
    (1, 'Risultati', []),
    (2, 'Validazione dei circuiti e sensibilità di Shor agli errori di porta', [
        'Le verifiche ideali dell’aritmetica confermano la correttezza dei circuiti senza '
        'rumore per [i]N[/i] = 15, [i]N[/i] = 21 e [i]N[/i] = 35; la distanza di variazione '
        'totale dalla distribuzione teorica è 1,38×10[sup]−9[/sup] per '
        '[i]N[/i] = 21 e 2,50×10[sup]−13[/sup] per [i]N[/i] = 35.',
        'Per ciascun valore di [i]p[/i][sub]g[/sub] si eseguono 20 repliche da 4.096 shot, '
        'misurando il successo per singolo shot. Senza errori il successo è del 74,89% '
        '(intervallo di confidenza al 95% [74,59%; 75,18%]), in accordo con il 75% atteso; '
        'scende al 72,5% per [i]p[/i][sub]g[/sub] = 0,1%, al 56,5% per 1%, al 45,2% per 2% e '
        'al 29,5% per 5%. Da 10% in poi resta intorno al 25%, fino al 24,59% per '
        '[i]p[/i][sub]g[/sub] = 0,5, praticamente il riferimento uniforme del 24,61%.',
        'La fragilità si spiega con il numero di porte: con 294 porte soggette a errore, '
        'già per [i]p[/i][sub]g[/sub] = 1% si verificano in media circa tre errori per '
        'esecuzione. Il risultato motiva le tre strategie successive.',
    ]),
    (2, 'Post-processing e machine learning', [
        'Negli scenari UC1 e UC2 le SVM selezionate raggiungono valori di F1 pari a 0,919 e '
        '0,923 e area sotto la curva ROC pari a 0,965 e 0,958, ma questa capacità '
        'predittiva non riduce il numero di iterazioni necessarie. TOP-4 ricava i fattori alla '
        'prima iterazione in entrambi gli scenari (1,00 iterazioni in media), contro 1,37 e '
        '1,50 di TOP-1, con differenze statisticamente significative nei test appaiati; il '
        'Metodo 2 richiede invece 1,50 e 1,37 iterazioni, senza migliorare significativamente '
        'TOP-1 e risultando peggiore di TOP-4. Il vantaggio viene quindi dal considerare più '
        'candidati, non dal classificatore; il risultato resta legato all’istanza, al '
        'rumore e ai metodi considerati.',
    ]),
    (2, 'Correzione quantistica degli errori', [
        'Le verifiche preliminari confermano il comportamento atteso: il codice a ripetizione '
        'riproduce, entro l’incertezza Monte Carlo, la relazione [i]p[/i][sub]L[/sub] = '
        '3[i]p[/i][sup]2[/sup] − 2[i]p[/i][sup]3[/sup], e per Steane la dipendenza '
        'dall’errore fisico è circa quadratica, con pendenza logaritmica 1,94 e '
        'pseudo-soglia intorno a 0,08. Nel surface code, sotto soglia l’aumento della '
        'distanza riduce l’errore logico e sopra soglia lo aumenta; il fit fornisce soglie '
        'di circa 0,86% in base [i]Z[/i] e 0,81% in base [i]X[/i].',
        'Con rumore nominale la rete neurale autonoma non supera significativamente MWPM. In '
        'presenza di crosstalk il decoder ibrido migliora invece le prestazioni: per '
        '[i]d[/i] = 3, [i]p[/i] = 0,003 e crosstalk pari a 0,01, [i]p[/i][sub]L[/sub] passa da '
        '0,05829 a 0,04851, una riduzione relativa del 16,8% significativa al test appaiato di '
        'McNemar; il beneficio si riduce a distanza 5 e non è significativo a distanza 7. '
        'Nelle configurazioni a distanza 5 BP+OSD supera l’ibrido basato su MWPM; le '
        'prestazioni dipendono inoltre dalla distanza, dal numero di cicli e dalla dimensione '
        'dell’ingresso della rete.',
        'Nel modello code-capacity il Gross code protegge 12 qubit logici con 144 qubit di '
        'dato meglio di 12 patch di surface code, cioè 12 blocchi indipendenti da un qubit '
        'logico ciascuno, a distanza 11 (1.452 qubit di dato) a ogni '
        'livello di rumore studiato, e meglio di 12 patch a distanza 13 (2.028) per '
        '[i]p[/i] ≥ 0,75%: a [i]p[/i] = 1% il blocco fallisce con probabilità '
        '1,0×10[sup]−6[/sup], contro 1,6×10[sup]−6[/sup]. Il risultato '
        'dipende dal decoder: con la schedulazione serial della propagazione delle credenze il '
        'Gross code corregge tutti gli errori fino a cinque qubit, verificati su circa '
        '5×10[sup]8[/sup] configurazioni, e a [i]p[/i] = 1% l’errore logico è circa '
        'cento volte più basso che con la configurazione parallela usata inizialmente, che '
        'falliva già con tre errori.',
    ]),
    (2, 'Riduzione del rumore con la AQFT', [
        'Nel pilota di Shor per [i]N[/i] = 15 l’AQFT selezionata aumenta il successo per '
        'shot dal 47,92% al 62,79%, cioè di 14,87 punti percentuali (+31% relativo) sui '
        'dati di verifica (intervallo di confidenza Newcombe al 95% [12,73; 16,99]), mentre '
        'le porte a due qubit ECR scendono da 362 a 279 (−22,9%) e la profondità da 1.351 '
        'a 1.141 (−15,5%). In questo caso il rumore evitato supera la precisione persa, e il '
        'miglioramento riguarda direttamente la fattorizzazione, non solo la stima di fase.',
        'Per [i]N[/i] = 21 l’AQFT applicata alle QFT interne all’aritmetica dimezza le '
        'porte ECR, da 49.661 a 23.178 (−53,3%), ma il successo ideale scende dal 46,67% al '
        '40,13% (−6,54 punti percentuali); nei tre confronti con rumore fra AQFT e QFT '
        'completa, con soli 128 shot di verifica per variante, gli intervalli delle differenze '
        'comprendono sempre lo zero, quindi per [i]N[/i] = 21 un beneficio dell’AQFT non '
        'è dimostrato. Nei nove benchmark di stima di fase tutte le differenze favoriscono '
        'l’AQFT, fino a 12,55 punti percentuali ([9,77; 15,30]). I confronti restano '
        'esplorativi e il grado di approssimazione dell’AQFT migliore fra quelli provati, '
        'due o tre, non vale in generale per registri di dimensione arbitraria.',
    ]),
    (1, 'Discussione e conclusioni', [
        'Lo studio ha mostrato che l’algoritmo di Shor è molto fragile al rumore: '
        'bastano pochi errori per degradarne il risultato. La selezione degli esiti riduce il '
        'numero di tentativi necessari, ma il vantaggio viene dal provare più candidati e '
        'non dal machine learning. La correzione quantistica degli errori protegge invece '
        'l’informazione, e un codice qLDPC come il Gross code ottiene la stessa protezione '
        'del surface code con molti meno qubit, purché si usi un decoder adatto. '
        'L’AQFT, infine, migliora il successo di Shor per [i]N[/i] = 15 riducendo le '
        'porte, mentre per [i]N[/i] = 21 un beneficio dell’AQFT non è dimostrato.',
        'Questi risultati valgono nel contesto simulativo considerato e non includono '
        'un’esecuzione fault-tolerant completa di Shor né una verifica su hardware '
        'reale.',
    ]),
    (1, 'Riferimenti bibliografici essenziali', []),
]

# Riferimenti nell'ordine di prima citazione nel testo Word (stesso testo di
# sezioni/05_bibliografia.tex; qui AlphaQubit precede Bravyi et al. perche' e' citato prima).
BIBLIOGRAFIA = [
    'P. W. Shor, “Algorithms for Quantum Computation: Discrete Logarithms and '
    'Factoring”, [i]Proc. 35th Annual Symposium on Foundations of Computer Science '
    '(FOCS)[/i], pp. 124–134, 1994.',
    'National Institute of Standards and Technology, [i]FIPS 203, FIPS 204 e FIPS 205[/i], '
    '2024.',
    'C. Gidney e M. Ekerå, “How to Factor 2048 Bit RSA Integers in 8 Hours Using '
    '20 Million Noisy Qubits”, [i]Quantum[/i], 5, 433, 2021.',
    'J. Preskill, “Quantum Computing in the NISQ Era and Beyond”, [i]Quantum[/i], '
    '2, 79, 2018.',
    'J. A. Smolin, G. Smith e A. Vargo, “Oversimplifying Quantum Factoring”, '
    '[i]Nature[/i], 499, pp. 163–165, 2013.',
    'S. Beauregard, “Circuit for Shor’s Algorithm Using 2[i]n[/i]+3 Qubits”, '
    '[i]Quantum Inf. Comput.[/i], 3(2), pp. 175–185, 2003.',
    'A. M. Steane, “Error Correcting Codes in Quantum Theory”, [i]Phys. Rev. '
    'Lett.[/i], 77(5), pp. 793–797, 1996.',
    'A. G. Fowler et al., “Surface Codes: Towards Practical Large-Scale Quantum '
    'Computation”, [i]Phys. Rev. A[/i], 86, 032324, 2012.',
    'Google Quantum AI and Collaborators, “Quantum Error Correction below the Surface '
    'Code Threshold”, [i]Nature[/i], 638, pp. 920–926, 2025.',
    'J. Bausch et al., “Learning High-Accuracy Error Decoding for Quantum '
    'Processors”, [i]Nature[/i], 635, pp. 834–840, 2024.',
    'S. Bravyi et al., “High-Threshold and Low-Overhead Fault-Tolerant Quantum '
    'Memory”, [i]Nature[/i], 627, pp. 778–782, 2024.',
    'J. Roffe et al., “Decoding Across the Quantum Low-Density Parity-Check Code '
    'Landscape”, [i]Phys. Rev. Research[/i], 2, 043423, 2020.',
    'A. Barenco, A. Ekert, K.-A. Suominen e P. Törmä, “Approximate Quantum '
    'Fourier Transform and Decoherence”, [i]Phys. Rev. A[/i], 54, pp. 139–146, '
    '1996.',
]


# ---------------------------------------------------------------------------
# Utilita' di formattazione
# ---------------------------------------------------------------------------
def _font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), FONT)


def _spaziatura(par, prima=0, dopo=3, interlinea=1.15):
    pf = par.paragraph_format
    pf.space_before = Pt(prima)
    pf.space_after = Pt(dopo)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = interlinea
    pf.first_line_indent = Cm(0)


def _bordo_sotto(elemento_pr, spessore_ottavi=4):
    """Filetto inferiore (spessore in ottavi di punto: 4 = 0,5 pt)."""
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(spessore_ottavi))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pbdr.append(bottom)
    elemento_pr.append(pbdr)


def _spazi_indivisibili(testo):
    """Evita che formule e citazioni si spezzino a fine riga, come il ~ di LaTeX."""
    testo = testo.replace(' = ', ' = ').replace(' ≡ ', ' ≡ ')
    testo = testo.replace(' \u2243 ', '\u00a0\u2243\u00a0').replace('mod ', 'mod\u00a0')
    testo = re.sub(r'(Metodo|Sezione) (\d)', '\\1\u00a0\\2', testo)
    return re.sub(r' \[(\d)', ' [\\1', testo)


def _senza_sillabazione(par):
    ppr = par._p.get_or_add_pPr()
    e = OxmlElement('w:suppressAutoHyphens')
    ppr.append(e)


def testo_con_markup(par, testo, size=12, bold=False):
    """Scrive il testo interpretando [i], [sup], [sub]."""
    testo = _spazi_indivisibili(testo)
    stato = {'i': False, 'sup': False, 'sub': False}
    for pezzo in re.split(r'(\[/?(?:i|sup|sub)\])', testo):
        m = re.fullmatch(r'\[(/?)(i|sup|sub)\]', pezzo)
        if m:
            stato[m.group(2)] = not m.group(1)
            continue
        if not pezzo:
            continue
        run = par.add_run(pezzo)
        _font(run, size=size, bold=bold, italic=stato['i'])
        if stato['sup']:
            run.font.superscript = True
        if stato['sub']:
            run.font.subscript = True


def _senza_bordi(tabella):
    tblpr = tabella._tbl.tblPr
    bordi = OxmlElement('w:tblBorders')
    for lato in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{lato}')
        e.set(qn('w:val'), 'nil')
        bordi.append(e)
    tblpr.append(bordi)


def _margini_cella_zero(tabella):
    tblpr = tabella._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for lato in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{lato}')
        e.set(qn('w:w'), '0')
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblpr.append(mar)


def _larghezze(tabella, cm):
    tabella.autofit = False
    for riga in tabella.rows:
        for cella, w in zip(riga.cells, cm):
            cella.width = Cm(w)


def _campo_pagina(par):
    """Inserisce il campo PAGE."""
    for tipo, testo in (('begin', None), (None, 'PAGE'), ('end', None)):
        run = par.add_run()
        _font(run, size=10)
        if tipo:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), tipo)
            run._element.append(fc)
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = testo
            run._element.append(it)


# ---------------------------------------------------------------------------
# Impaginazione
# ---------------------------------------------------------------------------
def imposta_pagina(sezione):
    sezione.page_width, sezione.page_height = Cm(21.0), Cm(29.7)
    sezione.top_margin = Cm(3)
    sezione.bottom_margin = Cm(3)
    sezione.left_margin = Cm(3.5)
    sezione.right_margin = Cm(3)
    # in main.tex: head 1,35 cm + headsep 0,45 cm sopra il corpo del testo
    sezione.header_distance = Cm(1.2)
    sezione.footer_distance = Cm(1.6)


def intestazione(sezione, testo_destra):
    header = sezione.header
    header.is_linked_to_previous = False
    par0 = header.paragraphs[0]
    tab = header.add_table(rows=1, cols=2, width=Cm(16.5))
    _senza_bordi(tab)
    _margini_cella_zero(tab)
    _larghezze(tab, [5.5, 11.0])
    sx, dx = tab.rows[0].cells
    p = sx.paragraphs[0]
    _spaziatura(p, dopo=0, interlinea=1.0)
    p.add_run().add_picture(LOGO, height=Cm(1.02))
    p = dx.paragraphs[0]
    _spaziatura(p, dopo=0, interlinea=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(testo_destra)
    _font(run, size=10, italic=True)
    # cella destra centrata in verticale, come \parbox[c]
    tcpr = dx._tc.get_or_add_tcPr()
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'center')
    tcpr.append(va)
    # il paragrafo vuoto iniziale dell'intestazione porta il filetto sotto la tabella
    par0._p.addnext(tab._tbl)
    header._element.append(par0._p)
    _spaziatura(par0, dopo=0, interlinea=1.0)
    _bordo_sotto(par0._p.get_or_add_pPr(), spessore_ottavi=3)
    run = par0.add_run()
    _font(run, size=2)


def pie_di_pagina(sezione):
    footer = sezione.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _campo_pagina(p)


def blocco_titolo(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spaziatura(p, prima=0, dopo=6)
    _senza_sillabazione(p)
    _font(p.add_run(TITOLO), size=14, bold=True)

    tab = doc.add_table(rows=2, cols=2)
    _senza_bordi(tab)
    _margini_cella_zero(tab)
    _larghezze(tab, [7.6, 8.9])
    tab.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r, (e1, v1, e2, v2) in enumerate(BLOCCO):
        for c, (etichetta, valore) in enumerate(((e1, v1), (e2, v2))):
            cp = tab.rows[r].cells[c].paragraphs[0]
            _spaziatura(cp, dopo=1, interlinea=1.0)
            _font(cp.add_run(f'{etichetta}: '), size=11, bold=True)
            _font(cp.add_run(valore), size=11)

    p = doc.add_paragraph()
    _spaziatura(p, prima=0, dopo=2, interlinea=1.0)
    _bordo_sotto(p._p.get_or_add_pPr(), spessore_ottavi=3)
    _font(p.add_run(), size=4)


def titolo(doc, numero, testo, livello):
    p = doc.add_paragraph()
    if livello == 1:
        _spaziatura(p, prima=9, dopo=2)
        size = 14
    else:
        _spaziatura(p, prima=6, dopo=1)
        size = 12
    p.paragraph_format.keep_with_next = True
    _senza_sillabazione(p)
    _font(p.add_run(f'{numero} {testo}'), size=size, bold=True)


def paragrafo(doc, testo):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _spaziatura(p)
    testo_con_markup(p, testo)


def _riduci_ultimo_paragrafo(doc):
    """Il paragrafo creato da un'interruzione di sezione e' vuoto: lo si rende alto 1 pt,
    cosi' non aggiunge righe bianche ne' pagine vuote."""
    p = doc.paragraphs[-1]
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(1)
    rpr = p._p.get_or_add_pPr()
    marca = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '2')
    marca.append(sz)
    rpr.append(marca)


def colonne(sezione, n, spazio_pt=10):
    sectpr = sezione._sectPr
    cols = sectpr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sectpr.append(cols)
    cols.set(qn('w:num'), str(n))
    cols.set(qn('w:space'), str(int(spazio_pt * 20)))


def main():
    doc = Document()
    stile = doc.styles['Normal']
    stile.font.name = FONT
    stile.font.size = Pt(12)
    stile.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    # lingua italiana e sillabazione automatica, come la sillabazione di LaTeX:
    # senza, il testo giustificato lascia spazi larghi fra le parole
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'it-IT')
    stile.element.rPr.append(lang)
    impostazioni = doc.settings.element
    sill = OxmlElement('w:autoHyphenation')
    sill.set(qn('w:val'), 'true')
    impostazioni.append(sill)

    sezione = doc.sections[0]
    imposta_pagina(sezione)
    intestazione(sezione, INTESTAZIONE)
    pie_di_pagina(sezione)

    blocco_titolo(doc)

    n1 = n2 = 0
    for livello, testo_titolo, paragrafi in CONTENUTO:
        if livello == 1:
            n1 += 1
            n2 = 0
            numero = str(n1)
        else:
            n2 += 1
            numero = f'{n1}.{n2}'
        if testo_titolo.startswith('Riferimenti') and BIBLIOGRAFIA:
            # come in 05_bibliografia.tex: titolo, poi due colonne separate da circa 10 pt,
            # corpo 6,4 pt, interlinea singola, testo a bandiera. Le colonne sono le due
            # celle di una tabella senza bordi: sempre bilanciate come in LaTeX e senza le
            # interruzioni di sezione, che in Word lasciavano una pagina bianca in fondo.
            titolo(doc, numero, testo_titolo, livello)
            tab = doc.add_table(rows=1, cols=3)
            _senza_bordi(tab)
            _margini_cella_zero(tab)
            _larghezze(tab, [7.15, 0.35, 7.15])
            meta = (len(BIBLIOGRAFIA) + 1) // 2
            for c, voci in ((0, BIBLIOGRAFIA[:meta]), (2, BIBLIOGRAFIA[meta:])):
                cella = tab.rows[0].cells[c]
                primo = c * 0 if c == 0 else meta
                for k, voce in enumerate(voci):
                    p = cella.paragraphs[0] if k == 0 else cella.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    _spaziatura(p, dopo=2, interlinea=1.0)
                    p.paragraph_format.left_indent = Cm(0.6)
                    p.paragraph_format.first_line_indent = Cm(-0.5)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.6))
                    _senza_sillabazione(p)
                    testo_con_markup(p, f'[{primo + k + 1}]\t{voce}', size=6.4)
            continue
        titolo(doc, numero, testo_titolo, livello)
        for t in paragrafi:
            paragrafo(doc, t)

    # Prima di sovrascrivere la copia OneDrive si controlla che nessuno l'abbia modificata
    # dopo l'ultima generazione: in quel caso ci si ferma e si mostrano le differenze, da
    # riportare in CONTENUTO prima di rigenerare.
    if os.path.exists(ONEDRIVE) and os.path.exists(USCITA):
        ultima, condivisa = _testo(USCITA), _testo(ONEDRIVE)
        if ultima != condivisa:
            print('STOP: la copia OneDrive e\' stata modificata dopo l\'ultima generazione.')
            for riga in difflib.unified_diff(ultima, condivisa, 'generata', 'onedrive',
                                             lineterm='', n=0):
                print(riga)
            raise SystemExit(1)

    doc.save(USCITA)
    print(f'Salvato: {USCITA}')
    shutil.copyfile(USCITA, ONEDRIVE)
    print(f'Copiato su OneDrive: {ONEDRIVE}')


def _testo(percorso):
    """Paragrafi del corpo del documento, per confrontare due versioni."""
    return [p.text for p in Document(percorso).paragraphs]


if __name__ == '__main__':
    main()
